"""
AI Resume ATS Checker
---------------------
Upload a resume (PDF / DOCX / TXT), optionally paste a job description, and get:
  * an estimated ATS score (0-100) with a category breakdown
  * strengths, issues, missing keywords and concrete improvements

UI: Streamlit   |   AI: Google Gemini Flash (google-genai SDK)
"""

import io
import json
import os
import re

import streamlit as st

# --------------------------------------------------------------------------
# Config
# --------------------------------------------------------------------------
DEFAULT_MODEL = "gemini-2.5-flash"   # change in the sidebar or via GEMINI_MODEL secret
MAX_FILE_MB = 5
MAX_RESUME_CHARS = 30_000            # safety limit sent to the model
MIN_TEXT_CHARS = 150                 # below this we assume a scanned / image PDF

CATEGORIES = {
    "formatting": "Formatting & Structure",
    "keywords": "Keywords & Skills",
    "experience": "Experience & Impact",
    "education": "Education & Certifications",
    "readability": "Readability & Length",
}


# --------------------------------------------------------------------------
# Text extraction
# --------------------------------------------------------------------------
def extract_text_from_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception:
        raise ValueError("Could not read this PDF. It may be corrupted; try re-exporting it.")
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            raise ValueError("This PDF is password-protected. Please upload an unlocked copy.")
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            pages.append("")
    return "\n".join(pages).strip()


def extract_text_from_docx(data: bytes) -> str:
    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
    except Exception:
        raise ValueError("Could not read this DOCX file. It may be corrupted or an old .doc file.")
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    # Resumes often keep content inside tables (two-column templates)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text.strip())
    return "\n".join(parts).strip()


def extract_text_from_txt(data: bytes) -> str:
    for enc in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(enc).strip()
        except UnicodeDecodeError:
            continue
    return ""


def extract_resume_text(filename: str, data: bytes) -> str:
    """Return plain text from an uploaded resume. Raises ValueError on bad input."""
    name = filename.lower()
    if name.endswith(".pdf"):
        return extract_text_from_pdf(data)
    if name.endswith(".docx"):
        return extract_text_from_docx(data)
    if name.endswith(".txt"):
        return extract_text_from_txt(data)
    raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")


# --------------------------------------------------------------------------
# Prompt + response handling
# --------------------------------------------------------------------------
SYSTEM_INSTRUCTION = (
    "You are an expert technical recruiter and ATS (Applicant Tracking System) "
    "specialist. You evaluate resumes strictly and honestly. You never invent "
    "facts that are not in the resume. You reply with valid JSON only."
)


def build_prompt(resume_text: str, job_description: str = "") -> str:
    jd = job_description.strip()
    if jd:
        jd_block = (
            "A target job description is provided. Score keyword match against it "
            "and list the important keywords from it that are missing from the resume.\n\n"
            f"<job_description>\n{jd[:10_000]}\n</job_description>\n"
        )
    else:
        jd_block = (
            "No job description was provided. Judge keywords against general "
            "industry standards for the role this resume appears to target.\n"
        )

    return f"""Evaluate the resume below for ATS compatibility and overall quality.

{jd_block}
<resume>
{resume_text[:MAX_RESUME_CHARS]}
</resume>

Scoring guidance (integers 0-100):
- formatting: clear section headings, consistent dates, contact info, no signs of tables/columns/graphics breaking parsing
- keywords: relevant hard skills, tools, and role keywords
- experience: action verbs, quantified achievements, relevance, progression
- education: education and certifications present and clearly listed
- readability: concise, well-organised, appropriate length (ideally 1-2 pages), no typos
- overall_score: weighted holistic ATS score, consistent with the category scores

Return ONLY a JSON object with exactly this shape:
{{
  "overall_score": <int 0-100>,
  "category_scores": {{
    "formatting": <int>, "keywords": <int>, "experience": <int>,
    "education": <int>, "readability": <int>
  }},
  "summary": "<2-3 sentence overall assessment>",
  "strengths": ["<short point>", ...],
  "issues": ["<problem found in this resume>", ...],
  "missing_keywords": ["<keyword>", ...],
  "improvements": [
    {{"priority": "High" | "Medium" | "Low", "section": "<resume section>",
      "suggestion": "<specific, actionable fix>",
      "example": "<optional rewritten example line, or empty string>"}}
  ]
}}
Give 3-6 strengths, 3-8 issues, up to 15 missing keywords and 5-10 improvements, ordered by priority."""


def _clamp_score(value, default=0) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError):
        return default


def _str_list(value) -> list:
    if not isinstance(value, list):
        return []
    return [str(v).strip() for v in value if str(v).strip()]


def parse_model_json(raw: str) -> dict:
    """Extract a JSON object from model output, tolerating ```json fences / extra text."""
    if not raw or not raw.strip():
        raise ValueError("The AI returned an empty response.")
    text = raw.strip()
    text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text, flags=re.IGNORECASE)
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        start, end = text.find("{"), text.rfind("}")
        if start != -1 and end > start:
            try:
                return json.loads(text[start : end + 1])
            except json.JSONDecodeError:
                pass
    raise ValueError("Could not read the AI response as JSON. Please try again.")


def normalize_result(data: dict) -> dict:
    """Validate / clean the model output so the UI never crashes on odd data."""
    if not isinstance(data, dict):
        raise ValueError("Unexpected response format from the AI.")

    raw_cats = data.get("category_scores") or {}
    cats = {k: _clamp_score(raw_cats.get(k)) for k in CATEGORIES}

    overall = data.get("overall_score")
    if overall is None:  # fall back to the mean of the categories
        overall = sum(cats.values()) / len(cats)

    improvements = []
    for item in data.get("improvements") or []:
        if not isinstance(item, dict):
            continue
        priority = str(item.get("priority", "Medium")).strip().capitalize()
        if priority not in ("High", "Medium", "Low"):
            priority = "Medium"
        suggestion = str(item.get("suggestion", "")).strip()
        if not suggestion:
            continue
        improvements.append(
            {
                "priority": priority,
                "section": str(item.get("section", "General")).strip() or "General",
                "suggestion": suggestion,
                "example": str(item.get("example", "") or "").strip(),
            }
        )
    order = {"High": 0, "Medium": 1, "Low": 2}
    improvements.sort(key=lambda i: order[i["priority"]])

    return {
        "overall_score": _clamp_score(overall),
        "category_scores": cats,
        "summary": str(data.get("summary", "")).strip(),
        "strengths": _str_list(data.get("strengths")),
        "issues": _str_list(data.get("issues")),
        "missing_keywords": _str_list(data.get("missing_keywords")),
        "improvements": improvements,
    }


# --------------------------------------------------------------------------
# Gemini call
# --------------------------------------------------------------------------
def analyze_resume(api_key: str, model: str, resume_text: str, job_description: str = "") -> dict:
    from google import genai
    from google.genai import types

    client = genai.Client(api_key=api_key)
    response = client.models.generate_content(
        model=model,
        contents=build_prompt(resume_text, job_description),
        config=types.GenerateContentConfig(
            system_instruction=SYSTEM_INSTRUCTION,
            temperature=0.2,
            response_mime_type="application/json",
        ),
    )
    return normalize_result(parse_model_json(response.text))


def friendly_error(exc: Exception) -> str:
    msg = str(exc)
    low = msg.lower()
    if "api key" in low or "api_key" in low or "permission" in low or "401" in low or "403" in low:
        return "Your Gemini API key looks invalid or lacks permission. Please check it."
    if "429" in low or "quota" in low or "rate" in low or "resource_exhausted" in low:
        return "Gemini rate limit / quota reached. Wait a minute and try again."
    if "404" in low or "not found" in low:
        return "That Gemini model name was not found. Change the model in the sidebar."
    return f"Something went wrong: {msg}"


# --------------------------------------------------------------------------
# UI helpers
# --------------------------------------------------------------------------
def score_label(score: int) -> str:
    if score >= 80:
        return "🟢 Excellent"
    if score >= 60:
        return "🟡 Good, but can improve"
    if score >= 40:
        return "🟠 Needs work"
    return "🔴 Poor"


def get_secret(name: str) -> str:
    """Read from Streamlit secrets, then environment; never crash if missing."""
    try:
        if name in st.secrets:
            return str(st.secrets[name])
    except Exception:
        pass
    return os.environ.get(name, "")


def render_results(result: dict) -> None:
    score = result["overall_score"]
    st.divider()
    col1, col2 = st.columns([1, 2])
    with col1:
        st.metric("Estimated ATS Score", f"{score} / 100")
        st.progress(score / 100)
        st.markdown(f"**{score_label(score)}**")
    with col2:
        st.subheader("Summary")
        st.write(result["summary"] or "No summary returned.")

    st.subheader("Score breakdown")
    for key, label in CATEGORIES.items():
        val = result["category_scores"][key]
        st.write(f"{label} — **{val}/100**")
        st.progress(val / 100)

    left, right = st.columns(2)
    with left:
        st.subheader("✅ Strengths")
        for s in result["strengths"] or ["None identified."]:
            st.markdown(f"- {s}")
    with right:
        st.subheader("⚠️ Issues found")
        for s in result["issues"] or ["None identified."]:
            st.markdown(f"- {s}")

    st.subheader("🔑 Missing keywords")
    if result["missing_keywords"]:
        st.write(", ".join(f"`{k}`" for k in result["missing_keywords"]))
    else:
        st.write("No major keyword gaps detected.")

    st.subheader("🛠️ Recommended improvements")
    icons = {"High": "🔴", "Medium": "🟠", "Low": "🟢"}
    for imp in result["improvements"]:
        with st.expander(f"{icons[imp['priority']]} {imp['priority']} — {imp['section']}"):
            st.write(imp["suggestion"])
            if imp["example"]:
                st.markdown("**Example:**")
                st.code(imp["example"], language=None)

    st.download_button(
        "⬇️ Download report (JSON)",
        data=json.dumps(result, indent=2),
        file_name="ats_report.json",
        mime="application/json",
    )


# --------------------------------------------------------------------------
# App
# --------------------------------------------------------------------------
def main() -> None:
    st.set_page_config(page_title="AI Resume ATS Checker", page_icon="📄", layout="wide")
    st.title("📄 AI Resume ATS Checker")
    st.caption("Upload your resume and get an estimated ATS score with actionable improvements.")

    with st.sidebar:
        st.header("Settings")
        api_key = get_secret("GEMINI_API_KEY")
        if api_key:
            st.success("API key loaded from secrets.")
        else:
            api_key = st.text_input(
                "Gemini API key",
                type="password",
                help="Get a free key at https://aistudio.google.com/app/apikey",
            )
        model = st.text_input("Gemini model", value=get_secret("GEMINI_MODEL") or DEFAULT_MODEL)
        st.info(
            "The score is an AI-based estimate, not the output of a real ATS. "
            "Use it as guidance. Resume text is sent to Google's Gemini API."
        )

    uploaded = st.file_uploader("Upload your resume", type=["pdf", "docx", "txt"])
    job_description = st.text_area(
        "Target job description (optional, but gives a better keyword match)",
        height=150,
        placeholder="Paste the job description here...",
    )

    if st.button("Analyze resume", type="primary"):
        if not api_key:
            st.error("Please enter your Gemini API key in the sidebar.")
            return
        if uploaded is None:
            st.error("Please upload a resume first.")
            return

        data = uploaded.getvalue()
        if len(data) > MAX_FILE_MB * 1024 * 1024:
            st.error(f"File is too large. Maximum size is {MAX_FILE_MB} MB.")
            return

        try:
            with st.spinner("Reading your resume..."):
                text = extract_resume_text(uploaded.name, data)
            if len(text) < MIN_TEXT_CHARS:
                st.error(
                    "Could not extract enough text. If your resume is a scanned image, "
                    "that is also a red flag for ATS systems — export a text-based PDF or DOCX."
                )
                return
            with st.spinner("Analyzing with Gemini..."):
                result = analyze_resume(api_key, model.strip() or DEFAULT_MODEL, text, job_description)
            st.session_state["result"] = result
        except ValueError as e:
            st.error(str(e))
            return
        except Exception as e:  # network / API errors
            st.error(friendly_error(e))
            return

    if "result" in st.session_state:
        render_results(st.session_state["result"])


if __name__ == "__main__":
    main()
